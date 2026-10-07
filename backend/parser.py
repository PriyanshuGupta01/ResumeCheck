"""In-memory resume parser for PDF and DOCX files.

Handles size limits, format checks, password protection detection,
and scanned document validation without writing to disk.
"""

import io
from typing import Tuple
from fastapi import HTTPException, status
import pdfplumber
import docx
from pdfminer.pdfdocument import PDFPasswordIncorrect

# Maximum allowed file size: 5 Megabytes
MAX_FILE_SIZE = 5 * 1024 * 1024


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from an uploaded PDF or DOCX resume in memory.

    Args:
        file_bytes: Binary content of the uploaded file.
        filename: Original filename to determine extension.

    Returns:
        Extracted plain text normalized with stripped whitespace.

    Raises:
        HTTPException: For invalid size, unsupported format, password protection,
                       or empty/scanned documents.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty. Please upload a valid resume.",
        )

    if len(file_bytes) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File size exceeds the 5 MB limit. Please upload a smaller document.",
        )

    lower_filename = filename.lower()

    if lower_filename.endswith(".pdf"):
        return _extract_from_pdf(file_bytes)
    elif lower_filename.endswith(".docx"):
        return _extract_from_docx(file_bytes)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a PDF (.pdf) or Word (.docx) file.",
        )


def _extract_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes using pdfplumber."""
    try:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            # Check if PDF is encrypted or password-protected
            if getattr(pdf.doc, "is_encrypted", False):
                # Try decrypting with empty password
                try:
                    pdf.doc.decrypt(b"")
                except Exception:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="This PDF is password-protected. Please upload an unprotected document.",
                    )

            for page_num, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())

        extracted_text = "\n\n".join(text_parts).strip()

        if not extracted_text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No readable text found. Scanned or image-only PDFs are not supported. Please upload a text-based document.",
            )

        return extracted_text

    except PDFPasswordIncorrect:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This PDF is password-protected. Please upload an unprotected document.",
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read PDF file: {str(e)}",
        )


def _extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            clean_line = paragraph.text.strip()
            if clean_line:
                text_parts.append(clean_line)

        # Extract text from tables (common in resume formatting)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate adjacent cells if merged
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    text_parts.append(" | ".join(unique_cells))

        extracted_text = "\n".join(text_parts).strip()

        if not extracted_text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The Word document does not contain any readable text. Please check the file and try again.",
            )

        return extracted_text

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not parse Word document: {str(e)}",
        )
